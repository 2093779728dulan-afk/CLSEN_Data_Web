"""Fail-closed support for the product-catalog live acceptance runner.

The PowerShell runner owns operating-system processes and the disposable MySQL
container.  This module owns database preflight, deterministic seed data,
workbook fixture mutation, and the final allowlisted evidence record.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
from dataclasses import asdict, dataclass, field, replace
from datetime import UTC, date, datetime, timedelta
from ipaddress import ip_address
from pathlib import Path
from typing import Any, Self
from urllib.parse import unquote
from uuid import uuid4

import sqlalchemy as sa
from app.auth.passwords import hash_password
from app.auth.sessions import Principal
from app.catalog.canonical_json import canonical_json_bytes
from app.catalog.return_dedup_repository import (
    FrozenPublicationMember,
    publication_set_digest,
)
from app.catalog.scope_repository import MySqlCatalogScopeRepository
from app.catalog.scope_service import CatalogScopeService, ConfirmScopeCommand
from app.catalog.snapshots import (
    MySqlCatalogSnapshotCoordinator,
    MySqlCatalogSnapshotWriter,
)
from app.catalog.workbook import parse_catalog_workbook
from app.db.engine import create_database_engine
from app.db.settings import DatabaseSettings
from app.imports.storage import LocalContentStore
from app.maintenance.schema_revision import CURRENT_SCHEMA_REVISION
from app.site_context import resolve_current_site_context
from openpyxl import load_workbook  # type: ignore[import-untyped]
from sqlalchemy.engine import Connection, Engine, make_url
from sqlalchemy.orm import Session

_RUN_ID = re.compile(r"[0-9a-f]{32}\Z", re.ASCII)
_DATABASE = re.compile(r"clsen_data_test_[0-9a-f]{16}\Z", re.ASCII)
_SHA256 = re.compile(r"[0-9a-f]{64}\Z", re.ASCII)
_LIVE_ROOT_PREFIX = "clsen-catalog-live-"
_DANGEROUS_SOURCE = 'LIVE-DANGER</script>&"RAW'


@dataclass(frozen=True, slots=True)
class DatabaseIdentity:
    database: str
    version: str
    version_comment: str
    character_set: str
    current_user: str
    tls_cipher: str


@dataclass(frozen=True, slots=True)
class ProcessIdentity:
    pid: int
    started_at: datetime
    command_sha256: str

    def __post_init__(self) -> None:
        if self.pid < 1:
            raise ValueError("process PID must be positive")
        if self.started_at.tzinfo is None or self.started_at.utcoffset() is None:
            raise ValueError("process start time must include a timezone")
        if _SHA256.fullmatch(self.command_sha256) is None:
            raise ValueError("process command digest must be SHA-256")


@dataclass(slots=True)
class RunManifest:
    run_id: str
    database_name: str
    runtime_root: str
    storage_root: str
    state: str = "created"
    details: dict[str, Any] = field(default_factory=dict)

    @classmethod
    def create(
        cls,
        *,
        run_id: str,
        database_name: str,
        runtime_root: Path,
        storage_root: Path,
    ) -> Self:
        validated = validate_run_id(run_id)
        expected_database = f"clsen_data_test_{validated[:16]}"
        if database_name != expected_database:
            raise ValueError("manifest database does not match the run ID")
        resolved_root = _assert_live_root(runtime_root, run_id=validated)
        resolved_storage = assert_owned_cleanup_path(
            storage_root,
            runtime_root=resolved_root,
            run_id=validated,
        )
        return cls(
            validated,
            database_name,
            str(resolved_root),
            str(resolved_storage),
        )

    @classmethod
    def from_dict(cls, value: dict[str, Any], *, expected_run_id: str) -> Self:
        required = {
            "run_id",
            "database_name",
            "runtime_root",
            "storage_root",
            "state",
        }
        if not required.issubset(value):
            raise ValueError("run manifest is incomplete")
        run_id = validate_run_id(str(value["run_id"]))
        if run_id != validate_run_id(expected_run_id):
            raise ValueError("run manifest run ID does not match its owner")
        manifest = cls.create(
            run_id=run_id,
            database_name=str(value["database_name"]),
            runtime_root=Path(str(value["runtime_root"])),
            storage_root=Path(str(value["storage_root"])),
        )
        manifest.state = str(value["state"])
        details = value.get("details", {})
        if not isinstance(details, dict):
            raise ValueError("run manifest details must be an object")
        manifest.details = details
        return manifest

    def require_state(self, expected: str) -> None:
        if self.state != expected:
            raise ValueError(
                f"run manifest state is {self.state!r}, expected {expected!r}"
            )

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def validate_run_id(value: str) -> str:
    if _RUN_ID.fullmatch(value) is None:
        raise ValueError("run ID must be 32 lower-case hexadecimal characters")
    return value


def _loopback_host(value: str | None) -> bool:
    if value is None:
        return False
    try:
        return ip_address(value).is_loopback
    except ValueError:
        return value.casefold() == "localhost"


def validate_database_urls(
    migration_value: str,
    runtime_value: str,
    *,
    run_id: str,
) -> tuple[DatabaseSettings, DatabaseSettings]:
    validated_run = validate_run_id(run_id)
    expected_database = f"clsen_data_test_{validated_run[:16]}"
    settings: list[DatabaseSettings] = []
    for label, value in (("migration", migration_value), ("runtime", runtime_value)):
        url = make_url(value)
        if url.drivername != "mysql+pymysql":
            raise ValueError(f"{label} URL must use mysql+pymysql")
        if not _loopback_host(url.host):
            raise ValueError(f"{label} URL must target loopback")
        if (
            url.database != expected_database
            or _DATABASE.fullmatch(url.database or "") is None
        ):
            raise ValueError(
                f"{label} URL must target the owned clsen_data_test schema"
            )
        if url.query.get("charset") != "utf8mb4":
            raise ValueError(f"{label} URL must require utf8mb4")
        if not url.query.get("ssl_ca"):
            raise ValueError(f"{label} URL must provide a verified TLS CA")
        settings.append(DatabaseSettings.from_url(value, pool_size=1, max_overflow=0))
    migration, runtime = settings
    if migration.url.username == runtime.url.username:
        raise ValueError("migration and runtime URLs must use separate accounts")
    if (
        migration.url.host,
        migration.url.port,
        migration.url.database,
    ) != (runtime.url.host, runtime.url.port, runtime.url.database):
        raise ValueError("migration and runtime URLs must target the same database")
    return migration, runtime


def assert_database_identity(
    identity: DatabaseIdentity,
    *,
    expected_database: str,
    expected_user: str,
) -> None:
    if (
        identity.database != expected_database
        or _DATABASE.fullmatch(identity.database) is None
    ):
        raise ValueError("database identity returned the wrong schema")
    if (
        not identity.version.startswith(("8.0.", "8.4."))
        or "mariadb" in (identity.version + identity.version_comment).casefold()
    ):
        raise ValueError("database must be Oracle MySQL 8.0 or 8.4")
    if "mysql" not in identity.version_comment.casefold():
        raise ValueError("database must be Oracle MySQL")
    if identity.character_set.casefold() != "utf8mb4":
        raise ValueError("database schema must use utf8mb4")
    if identity.current_user.partition("@")[0] != expected_user:
        raise ValueError("database authenticated account does not match its URL")
    if not identity.tls_cipher.strip():
        raise ValueError("database connection must negotiate TLS")


def assert_empty_seed_target(counts: dict[str, int]) -> None:
    expected = {"users", "facts", "catalog"}
    if set(counts) != expected or any(
        type(value) is not int or value < 0 for value in counts.values()
    ):
        raise ValueError("seed target counts are invalid")
    if any(counts.values()):
        raise ValueError(
            "refusing to seed a database containing ordinary users or facts"
        )


def _assert_live_root(path: Path, *, run_id: str) -> Path:
    resolved = path.resolve()
    if resolved.name != f"{_LIVE_ROOT_PREFIX}{validate_run_id(run_id)}":
        raise ValueError("runtime root is not named for the owned run ID")
    if resolved == Path(resolved.anchor):
        raise ValueError("runtime root must not be a filesystem root")
    return resolved


def assert_owned_cleanup_path(
    path: Path,
    *,
    runtime_root: Path,
    run_id: str,
) -> Path:
    root = _assert_live_root(runtime_root, run_id=run_id)
    resolved = path.resolve()
    if resolved == root or root not in resolved.parents:
        raise ValueError("cleanup target is outside the runner-owned runtime root")
    return resolved


def assert_process_identity(
    expected: ProcessIdentity,
    actual: ProcessIdentity,
) -> None:
    if expected.pid != actual.pid:
        raise ValueError("process PID does not match")
    if abs((expected.started_at - actual.started_at).total_seconds()) > 0.001:
        raise ValueError("process start time changed; PID may have been reused")
    if expected.command_sha256 != actual.command_sha256:
        raise ValueError("process command fingerprint changed")


def assert_live_playwright_config(source: str) -> None:
    if re.search(r"\bwebServer\s*:", source) or re.search(
        r"\breuseExistingServer\s*:", source
    ):
        raise ValueError("live Playwright config must not start or reuse a server")


def _manifest_path(value: str, *, run_id: str) -> Path:
    path = Path(value).resolve()
    root = _assert_live_root(path.parent, run_id=run_id)
    assert_owned_cleanup_path(path, runtime_root=root, run_id=run_id)
    return path


def _load_manifest(path: Path, *, run_id: str) -> RunManifest:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError("run manifest must be a JSON object")
    return RunManifest.from_dict(value, expected_run_id=run_id)


def _save_manifest(path: Path, manifest: RunManifest) -> None:
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(
        json.dumps(manifest.to_dict(), ensure_ascii=False, indent=2, sort_keys=True),
        encoding="utf-8",
    )
    temporary.replace(path)


def _database_identity(engine: Engine) -> DatabaseIdentity:
    with engine.connect() as connection:
        row = (
            connection.execute(
                sa.text(
                    "SELECT DATABASE() AS db, VERSION() AS version, "
                    "@@version_comment AS version_comment, "
                    "@@character_set_database AS character_set, "
                    "CURRENT_USER() AS authenticated_user"
                )
            )
            .mappings()
            .one()
        )
        tls = connection.execute(sa.text("SHOW SESSION STATUS LIKE 'Ssl_cipher'")).one()
    return DatabaseIdentity(
        database=str(row["db"]),
        version=str(row["version"]),
        version_comment=str(row["version_comment"]),
        character_set=str(row["character_set"]),
        current_user=str(row["authenticated_user"]),
        tls_cipher=str(tls[1]),
    )


def _seed_target_counts(engine: Engine) -> dict[str, int]:
    with engine.connect() as connection:
        row = (
            connection.execute(
                sa.text(
                    "SELECT (SELECT COUNT(*) FROM users) AS users, "
                    "(SELECT COUNT(*) FROM daily_datasets) AS facts, "
                    "((SELECT COUNT(*) FROM product_catalog_versions) + "
                    " (SELECT COUNT(*) FROM catalog_exports)) AS catalog"
                )
            )
            .mappings()
            .one()
        )
    return {name: int(row[name]) for name in ("users", "facts", "catalog")}


def _settings(run_id: str) -> tuple[DatabaseSettings, DatabaseSettings]:
    migration_value = os.environ.get("APP_MIGRATION_DATABASE_URL", "")
    runtime_value = os.environ.get("APP_TEST_RUNTIME_DATABASE_URL", "")
    if not migration_value or not runtime_value:
        raise ValueError("live database URLs are required in the process environment")
    return validate_database_urls(migration_value, runtime_value, run_id=run_id)


def preflight(*, run_id: str, manifest_path: Path) -> None:
    manifest = _load_manifest(manifest_path, run_id=run_id)
    manifest.require_state("created")
    migration_settings, runtime_settings = _settings(run_id)
    migration = create_database_engine(migration_settings)
    runtime = create_database_engine(runtime_settings)
    try:
        assert_database_identity(
            _database_identity(migration),
            expected_database=manifest.database_name,
            expected_user=str(migration_settings.url.username),
        )
        assert_database_identity(
            _database_identity(runtime),
            expected_database=manifest.database_name,
            expected_user=str(runtime_settings.url.username),
        )
        assert_empty_seed_target(_seed_target_counts(migration))
        with migration.connect() as connection:
            revision = str(
                connection.scalar(sa.text("SELECT version_num FROM alembic_version"))
            )
        if revision != CURRENT_SCHEMA_REVISION:
            raise ValueError("live database is not at the current migration head")
    finally:
        runtime.dispose()
        migration.dispose()
    manifest.state = "preflighted"
    manifest.details["migration_head"] = CURRENT_SCHEMA_REVISION
    _save_manifest(manifest_path, manifest)


def _insert_admin(engine: Engine, *, username: str, password: str) -> int:
    with engine.begin() as connection:
        inserted = connection.execute(
            sa.text(
                "INSERT INTO users (username_normalized, display_name, "
                "password_hash, status, created_at, updated_at) "
                "VALUES (:username, 'Catalog Live Admin', :password_hash, "
                "'active', UTC_TIMESTAMP(6), UTC_TIMESTAMP(6))"
            ),
            {"username": username.casefold(), "password_hash": hash_password(password)},
        )
        if inserted.lastrowid is None:
            raise RuntimeError("live seed did not return an administrator ID")
        user_id = int(inserted.lastrowid)
        connection.execute(
            sa.text(
                "INSERT INTO user_roles (user_id, role_id, assigned_by, assigned_at) "
                "SELECT :user_id, id, :user_id, UTC_TIMESTAMP(6) "
                "FROM roles WHERE code='admin'"
            ),
            {"user_id": user_id},
        )
    return user_id


def _confirm_scope(engine: Engine, *, admin_id: int, run_id: str) -> int:
    principal = Principal(
        user_id=admin_id,
        username=f"catalog-live-admin-{run_id}",
        display_name="Catalog Live Admin",
        roles=frozenset({"admin"}),
        session_id=admin_id,
        csrf_token_hash=b"x" * 32,
    )
    with Session(engine) as session:
        context = resolve_current_site_context(session, principal)
    result = CatalogScopeService(MySqlCatalogScopeRepository(engine)).confirm(
        context,
        ConfirmScopeCommand(
            operation_token=uuid4(),
            expected_site_config_revision=context.site_config_revision,
            seller_account_scope_key=f"catalog-live-{run_id}",
            seller_account_display_name="Catalog Live Seller",
            marketplace_scope_code="amazon_us",
            reason="Confirm runner-owned catalog live source scope",
        ),
    )
    if result.report_source_scope_id is None:
        raise RuntimeError("live scope confirmation did not return a scope ID")
    return result.report_source_scope_id


def _seed_ready_dedup(
    engine: Engine,
    *,
    scope_id: int,
    return_dataset_id: int,
    return_version_id: int,
    report_date: date,
) -> int:
    with engine.connect() as connection:
        data_revision = int(
            connection.scalar(
                sa.text("SELECT data_revision FROM dataset_versions WHERE id=:id"),
                {"id": return_version_id},
            )
        )
    member = FrozenPublicationMember(
        dataset_id=return_dataset_id,
        published_version_id=return_version_id,
        publication_revision=0,
        version_data_revision=data_revision,
        report_date=report_date,
    )
    set_digest = publication_set_digest((member,))
    decision_digest = hashlib.sha256(canonical_json_bytes([])).digest()
    with engine.begin() as connection:
        inserted = connection.execute(
            sa.text(
                "INSERT INTO return_dedup_snapshots (site_id, "
                "report_source_scope_id, return_publication_set_digest, "
                "rule_version, base_snapshot_id, checkpoint_snapshot_id, "
                "delta_depth, snapshot_kind, status, decision_digest) VALUES "
                "(1, :scope, :digest, 'return_event_fingerprint_v1', NULL, "
                "NULL, 0, 'checkpoint', 'ready', :decision)"
            ),
            {"scope": scope_id, "digest": set_digest, "decision": decision_digest},
        )
        if inserted.lastrowid is None:
            raise RuntimeError("live seed did not return a dedup snapshot ID")
        snapshot_id = int(inserted.lastrowid)
        updated = connection.execute(
            sa.text(
                "UPDATE return_dedup_publications SET current_snapshot_id=:snapshot, "
                "dedup_publication_revision=1 WHERE site_id=1 "
                "AND report_source_scope_id=:scope"
            ),
            {"snapshot": snapshot_id, "scope": scope_id},
        )
        if updated.rowcount != 1:
            raise RuntimeError("live seed could not bind the dedup publication")
    return snapshot_id


def _expand_business_fixture(
    connection: Connection,
    *,
    version_id: int,
) -> None:
    metadata = sa.MetaData()
    version_rows_table = sa.Table("version_rows", metadata, autoload_with=connection)
    business_rows_table = sa.Table(
        "business_report_rows", metadata, autoload_with=connection
    )
    source_row_value = connection.scalar(
        sa.select(business_rows_table.c.version_row_id)
        .where(business_rows_table.c.version_id == version_id)
        .order_by(business_rows_table.c.version_row_id)
        .limit(1)
    )
    if source_row_value is None:
        raise RuntimeError("live business fixture has no template row")
    source_row_id = int(source_row_value)
    version_template = dict(
        connection.execute(
            sa.select(version_rows_table).where(
                version_rows_table.c.id == source_row_id
            )
        )
        .mappings()
        .one()
    )
    business_template = dict(
        connection.execute(
            sa.select(business_rows_table).where(
                business_rows_table.c.version_row_id == source_row_id
            )
        )
        .mappings()
        .one()
    )
    next_ordinal = int(
        connection.scalar(
            sa.select(sa.func.max(version_rows_table.c.row_ordinal)).where(
                version_rows_table.c.version_id == version_id
            )
        )
        or 0
    )
    additional_variants = (
        "100-RED-S",
        "100-RED-M",
        "100-RED-L",
        "100-RED-XL",
        "100-RED-2XL",
        "100-RED-3XL",
        "100-RED-4XL",
        "100-BLUE-5XL",
        "100-GREEN-S",
        "200-YELLOW-XXL",
        "EXCLUDE-ME",
    )
    for offset, seller_sku in enumerate(additional_variants, start=1):
        version_values = {
            key: value for key, value in version_template.items() if key != "id"
        }
        version_values.update(
            {
                "logical_row_id": uuid4().bytes,
                "row_ordinal": next_ordinal + offset,
            }
        )
        inserted = connection.execute(
            version_rows_table.insert().values(version_values)
        )
        if inserted.lastrowid is None:
            raise RuntimeError("live scale fixture did not create a version row")
        business_values = {
            key: value
            for key, value in business_template.items()
            if key != "version_row_id"
        }
        business_values.update(
            {
                "version_row_id": int(inserted.lastrowid),
                "sku": seller_sku,
                "title": f"Synthetic evidence title {offset % 3}",
                "sessions_total": 0 if offset == 1 else offset * 10,
                "units_ordered": 0 if offset == 1 else offset,
                "ordered_product_sales": offset * 10,
            }
        )
        connection.execute(business_rows_table.insert().values(business_values))


def seed(
    *,
    run_id: str,
    manifest_path: Path,
    include_catalog_experience: bool = False,
) -> None:
    from tests.integration.dataset_test_support import DatasetIntegrationCase
    from tests.integration.publication_test_support import PublicationIntegrationCase

    manifest = _load_manifest(manifest_path, run_id=run_id)
    manifest.require_state("preflighted")
    password = os.environ.get("CATALOG_LIVE_ADMIN_PASSWORD", "")
    username = os.environ.get("CATALOG_LIVE_ADMIN_USERNAME", "")
    if len(password) < 12 or username != f"catalog-live-admin-{run_id}":
        raise ValueError("live administrator credentials are missing or not run-scoped")
    migration_settings, runtime_settings = _settings(run_id)
    migration = create_database_engine(migration_settings)
    manifest.state = "seeding"
    _save_manifest(manifest_path, manifest)
    try:
        assert_empty_seed_target(_seed_target_counts(migration))
        admin_id = _insert_admin(migration, username=username, password=password)
        storage_root = Path(manifest.storage_root)
        storage_root.mkdir(parents=True, exist_ok=False)
        dataset = DatasetIntegrationCase(
            migration_engine=migration,
            runtime_settings=runtime_settings,
            store=LocalContentStore(storage_root, max_upload_bytes=16 * 1024 * 1024),
            owner_user_id=admin_id,
            other_manager_user_id=admin_id,
            admin_user_id=admin_id,
            user_ids=(admin_id,),
        )
        publication = PublicationIntegrationCase(dataset)
        report_date = date(2025, 7, 17)
        if include_catalog_experience:
            # Keep interactive sample facts inside the existing 366-day query limit.
            report_date = (date.today().replace(day=1) - timedelta(days=1)).replace(
                day=17
            )
        _business_baseline, business = publication.materialize_baseline(
            report_date=report_date,
            variant="live-business",
            report_type="business_report",
        )
        _return_baseline, returns = publication.materialize_baseline(
            report_date=report_date,
            variant="live-returns",
            report_type="return_report",
        )
        _pending_baseline, pending_business = publication.materialize_baseline(
            report_date=report_date.replace(day=18),
            variant="live-pending-business",
            report_type="business_report",
        )
        with migration.begin() as connection:
            business_rows = list(
                connection.scalars(
                    sa.text(
                        "SELECT version_row_id FROM business_report_rows "
                        "WHERE version_id=:version ORDER BY version_row_id"
                    ),
                    {"version": business.version_id},
                )
            )
            if len(business_rows) != 2:
                raise RuntimeError("live business fixture did not contain two rows")
            for row_id, sku, title in (
                (int(business_rows[0]), "100-RED-XS", _DANGEROUS_SOURCE),
                (int(business_rows[1]), "OLD100-GREEN-S", "Legacy Seller SKU"),
            ):
                connection.execute(
                    sa.text(
                        "UPDATE business_report_rows SET sku=:sku, title=:title "
                        "WHERE version_row_id=:row_id AND version_id=:version"
                    ),
                    {
                        "sku": sku,
                        "title": title,
                        "row_id": row_id,
                        "version": business.version_id,
                    },
                )
            _expand_business_fixture(
                connection,
                version_id=business.version_id,
            )
            return_rows = list(
                connection.scalars(
                    sa.text(
                        "SELECT version_row_id FROM return_report_rows "
                        "WHERE version_id=:version ORDER BY version_row_id"
                    ),
                    {"version": returns.version_id},
                )
            )
            for row_id, sku in zip(
                return_rows,
                ("100-RED-XS", "100-GREEN-S"),
                strict=False,
            ):
                connection.execute(
                    sa.text(
                        "UPDATE return_report_rows SET sku=:sku "
                        "WHERE version_row_id=:row_id AND version_id=:version"
                    ),
                    {
                        "sku": sku,
                        "row_id": int(row_id),
                        "version": returns.version_id,
                    },
                )
            pending_rows = list(
                connection.scalars(
                    sa.text(
                        "SELECT version_row_id FROM business_report_rows "
                        "WHERE version_id=:version ORDER BY version_row_id"
                    ),
                    {"version": pending_business.version_id},
                )
            )
            if len(pending_rows) != 2:
                raise RuntimeError("live pending fixture did not contain two rows")
            for row_id, sku, title in (
                (int(pending_rows[0]), "100-RED-XS", "Existing red variant"),
                (int(pending_rows[1]), "100-NAVY-L", "Pending navy large variant"),
            ):
                connection.execute(
                    sa.text(
                        "UPDATE business_report_rows SET sku=:sku, title=:title "
                        "WHERE version_row_id=:row_id AND version_id=:version"
                    ),
                    {
                        "sku": sku,
                        "title": title,
                        "row_id": row_id,
                        "version": pending_business.version_id,
                    },
                )
        dataset.mark_published(business.dataset_id, business.version_id)
        dataset.mark_published(returns.dataset_id, returns.version_id)
        runtime = dataset.runtime_engine()
        try:
            scope_id = _confirm_scope(runtime, admin_id=admin_id, run_id=run_id)
        finally:
            runtime.dispose()
        dedup_id = _seed_ready_dedup(
            migration,
            scope_id=scope_id,
            return_dataset_id=returns.dataset_id,
            return_version_id=returns.version_id,
            report_date=report_date,
        )
        manifest.details.update(
            {
                "admin_user_id": admin_id,
                "business_dataset_id": business.dataset_id,
                "return_dataset_id": returns.dataset_id,
                "source_scope_id": scope_id,
                "dedup_snapshot_id": dedup_id,
            }
        )
        if not include_catalog_experience:
            manifest.state = "seeded"
            _save_manifest(manifest_path, manifest)
            return
        from app.catalog.export_service import MySqlCatalogExportFinalizer
        from app.catalog.import_repository import MySqlCatalogImportRepository
        from app.catalog.import_service import project_catalog_workbook
        from app.catalog.queries import MySqlCatalogQueries
        from app.catalog.service import CatalogService
        from app.imports.storage import StoredContent
        from app.worker.catalog_exports import MySqlCatalogExportWorkSource

        from tests.integration.test_catalog_repository import _completed_projection

        catalog_runtime = dataset.runtime_engine()
        try:
            card_frozen = MySqlCatalogSnapshotCoordinator(catalog_runtime).freeze(
                site_id=1
            )
            card_export = MySqlCatalogSnapshotWriter(catalog_runtime).persist_building(
                card_frozen,
                operation_token=uuid4(),
                created_by_user_id=admin_id,
            )
            catalog_export = MySqlCatalogExportWorkSource(
                catalog_runtime
            ).next_building()
            if (
                catalog_export is None
                or catalog_export.export_id != card_export.export_id
            ):
                raise RuntimeError("live catalog card export could not be claimed")
            draft = project_catalog_workbook(catalog_export.workbook)
            MySqlCatalogExportFinalizer(catalog_runtime).finalize_ready(
                catalog_export,
                artifact=StoredContent(
                    "sha256/cc/" + "c" * 64,
                    "c" * 64,
                    123,
                    False,
                ),
                capacity_violations=(),
            )
            projection = _completed_projection(draft)
            mapped_rows = [
                row
                for row in projection.rows
                if row.section == "mappings"
                and isinstance(row.payload.get("target_sku_row_key"), str)
            ]
            if len(mapped_rows) < 2:
                raise RuntimeError("live catalog card fixture needs two Seller SKUs")
            shared_target = mapped_rows[0].payload["target_sku_row_key"]
            alias_row_key = mapped_rows[1].row_key
            card_rows = []
            for row in projection.rows:
                if row.row_key != alias_row_key:
                    card_rows.append(row)
                    continue
                payload = dict(row.payload)
                payload["target_sku_row_key"] = shared_target
                payload["identity_decision"] = "mapped"
                payload["exclusion_reason"] = None
                card_rows.append(
                    replace(
                        row,
                        payload=payload,
                        payload_digest=hashlib.sha256(
                            canonical_json_bytes(payload)
                        ).digest(),
                    )
                )
            projection = replace(projection, rows=tuple(card_rows))
            imports = MySqlCatalogImportRepository(catalog_runtime)
            catalog_import_id = imports.create_online_from_export(
                site_id=1,
                owner_user_id=admin_id,
                operation_token=uuid4(),
                export_id=catalog_export.export_id,
                expected_catalog_publication_revision=0,
                projection=projection,
            )
            imports.validate_import(
                import_id=catalog_import_id,
                site_id=1,
                actor_user_id=admin_id,
                operation_token=uuid4(),
                expected_workflow_revision=0,
            )
            imports.transition_import(
                import_id=catalog_import_id,
                site_id=1,
                actor_user_id=admin_id,
                operation_token=uuid4(),
                expected_workflow_revision=1,
                command="submit",
            )
            service = CatalogService(catalog_runtime)
            preview = service.preview_import(
                import_id=catalog_import_id,
                site_id=1,
                actor_user_id=admin_id,
                actor_roles=frozenset({"admin"}),
                operation_token=uuid4(),
                expected_workflow_revision=2,
                now=datetime(2026, 9, 7, 12, 0, tzinfo=UTC),
            )
            approved = service.approve_import(
                import_id=catalog_import_id,
                site_id=1,
                actor_user_id=admin_id,
                actor_roles=frozenset({"admin"}),
                operation_token=uuid4(),
                expected_workflow_revision=2,
                expected_catalog_publication_revision=0,
                projection_id=preview.projection_id,
                projection_digest=preview.projection_digest,
                reason="Seed the isolated catalog card experience",
                now=datetime(2026, 9, 7, 12, 1, tzinfo=UTC),
            )
            current_catalog = MySqlCatalogQueries(catalog_runtime).current_catalog(
                site_id=1
            )
            if current_catalog is None:
                raise RuntimeError("live catalog card approval was not published")
            pending_style = next(
                (
                    style
                    for style in current_catalog.styles
                    if style.style_code == "100"
                ),
                None,
            )
            if pending_style is None:
                raise RuntimeError("live pending fixture could not resolve style 100")
            dataset.mark_published(
                pending_business.dataset_id,
                pending_business.version_id,
            )
            pending_frozen = MySqlCatalogSnapshotCoordinator(catalog_runtime).freeze(
                site_id=1
            )
            pending_export = MySqlCatalogSnapshotWriter(
                catalog_runtime
            ).persist_building(
                pending_frozen,
                operation_token=uuid4(),
                created_by_user_id=admin_id,
            )
            pending_command = MySqlCatalogExportWorkSource(
                catalog_runtime
            ).next_building()
            if (
                pending_command is None
                or pending_command.export_id != pending_export.export_id
            ):
                raise RuntimeError("live pending export could not be claimed")
            pending_projection = project_catalog_workbook(pending_command.workbook)
            MySqlCatalogExportFinalizer(catalog_runtime).finalize_ready(
                pending_command,
                artifact=StoredContent(
                    "sha256/dd/" + "d" * 64,
                    "d" * 64,
                    124,
                    False,
                ),
                capacity_violations=(),
            )
            pending_import_id = imports.create_online_from_export(
                site_id=1,
                owner_user_id=admin_id,
                operation_token=uuid4(),
                export_id=pending_export.export_id,
                expected_catalog_publication_revision=1,
                projection=pending_projection,
            )
            pending_items = MySqlCatalogQueries(catalog_runtime).style_pending_items(
                site_id=1,
                style_id=pending_style.style_id,
            )
            from app.catalog.identity_conflicts import identity_owners

            for pending_row in pending_projection.rows:
                if pending_row.section == "mappings":
                    identity_owners(
                        catalog_runtime,
                        site_id=1,
                        import_id=pending_import_id,
                        row_key=pending_row.row_key,
                    )
            matching_pending_items = [
                item for item in pending_items if item["import_id"] == pending_import_id
            ]
            if len(matching_pending_items) != 1:
                raise RuntimeError("live pending import is not visible from its style")
            pending_item = matching_pending_items[0]
            if (
                pending_item["variant_count"] != 1
                or pending_item["seller_sku_count"] != 1
                or pending_item["seller_sku_previews"] != ("100-NAVY-L",)
            ):
                raise RuntimeError("live pending import summary is invalid")
        finally:
            catalog_runtime.dispose()
        runtime = dataset.runtime_engine()
        try:
            frozen = MySqlCatalogSnapshotCoordinator(runtime).freeze(site_id=1)
            recovery_export = MySqlCatalogSnapshotWriter(runtime).persist_building(
                frozen,
                operation_token=uuid4(),
                created_by_user_id=admin_id,
            )
            with runtime.begin() as connection:
                updated = connection.execute(
                    sa.text(
                        """
                        UPDATE catalog_jobs
                        SET status = 'failed', job_revision = job_revision + 1,
                            attempt_count = 3,
                            execution_phase = 'awaiting_manual_recovery',
                            retry_allowed = TRUE,
                            failure_reason_code = 'catalog_export_worker_error'
                        WHERE job_kind = 'catalog_export'
                          AND resource_kind = 'catalog_export'
                          AND resource_id = :export_id
                          AND status = 'queued'
                        """
                    ),
                    {"export_id": recovery_export.export_id},
                )
                if updated.rowcount != 1:
                    raise RuntimeError("live recovery export was not seeded")
        finally:
            runtime.dispose()
        manifest.details.update(
            {
                "admin_user_id": admin_id,
                "business_dataset_id": business.dataset_id,
                "return_dataset_id": returns.dataset_id,
                "source_scope_id": scope_id,
                "dedup_snapshot_id": dedup_id,
                "catalog_version_id": approved.catalog_version_id,
                "pending_style_id": pending_style.style_id,
                "pending_import_id": pending_import_id,
                "recovery_export_id": recovery_export.export_id,
            }
        )
        manifest.state = "seeded"
        _save_manifest(manifest_path, manifest)
    except Exception:
        manifest.state = "seed_failed"
        _save_manifest(manifest_path, manifest)
        raise
    finally:
        migration.dispose()


def mutate_workbook(
    *,
    run_id: str,
    manifest_path: Path,
    input_path: Path,
    output_path: Path,
) -> None:
    manifest = _load_manifest(manifest_path, run_id=run_id)
    if manifest.state not in {"seeded", "running"}:
        raise ValueError("workbook mutation requires a running live manifest")
    root = Path(manifest.runtime_root)
    source = assert_owned_cleanup_path(input_path, runtime_root=root, run_id=run_id)
    target = assert_owned_cleanup_path(output_path, runtime_root=root, run_id=run_id)
    if not source.is_file() or target.exists():
        raise ValueError("workbook fixture paths are not in the expected state")
    workbook = load_workbook(source, read_only=False, data_only=False, keep_links=False)
    try:
        if tuple(workbook.sheetnames) != (
            "生成信息",
            "款式主档",
            "SKU变体",
            "SellerSKU映射",
            "待处理异常",
            "来源证据",
        ):
            raise ValueError("downloaded catalog workbook has unexpected sheets")
        cell = workbook["SKU变体"]["E2"]
        if not isinstance(cell.value, str) or not cell.value:
            raise ValueError("downloaded catalog workbook has no editable E field")
        cell.value = "Live Blue E-field"
        workbook.save(target)
    finally:
        workbook.close()
    content = target.read_bytes()
    parse_catalog_workbook(content)
    manifest.details["mutated_workbook_sha256"] = hashlib.sha256(content).hexdigest()
    _save_manifest(manifest_path, manifest)


def _read_owned_json(path: Path, *, manifest: RunManifest) -> dict[str, Any]:
    owned = assert_owned_cleanup_path(
        path,
        runtime_root=Path(manifest.runtime_root),
        run_id=manifest.run_id,
    )
    value = json.loads(owned.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError("live result must be a JSON object")
    return value


def collect_evidence(
    *,
    run_id: str,
    manifest_path: Path,
    result_path: Path,
    evidence_path: Path,
) -> dict[str, Any]:
    manifest = _load_manifest(manifest_path, run_id=run_id)
    if manifest.state not in {"seeded", "running"}:
        raise ValueError("evidence requires a completed live execution")
    result = _read_owned_json(result_path, manifest=manifest)
    online_import_id = int(result["online_import_id"])
    raw_import_id = int(result["raw_import_id"])
    expected_version_id = int(result["catalog_version_id"])
    downloaded = assert_owned_cleanup_path(
        Path(str(result["download_path"])),
        runtime_root=Path(manifest.runtime_root),
        run_id=run_id,
    )
    workbook_checksum = hashlib.sha256(downloaded.read_bytes()).hexdigest()
    if workbook_checksum != str(result["download_sha256"]):
        raise ValueError("downloaded workbook checksum does not match the live result")
    migration_settings, _runtime_settings = _settings(run_id)
    engine = create_database_engine(migration_settings)
    try:
        with engine.connect() as connection:
            row = (
                connection.execute(
                    sa.text(
                        "SELECT publications.current_catalog_version_id AS current_id, "
                        "versions.version_number, versions.source_import_id, "
                        "imports.export_id, exports.artifact_sha256, "
                        "imports.workflow_revision "
                        "FROM product_catalog_publications AS publications "
                        "INNER JOIN product_catalog_versions AS versions "
                        "ON versions.id=publications.current_catalog_version_id "
                        "INNER JOIN catalog_imports AS imports "
                        "ON imports.id=versions.source_import_id "
                        "INNER JOIN catalog_exports AS exports "
                        "ON exports.id=imports.export_id "
                        "WHERE publications.site_id=1"
                    )
                )
                .mappings()
                .one()
            )
            revision = str(
                connection.scalar(sa.text("SELECT version_num FROM alembic_version"))
            )
            raw_status = str(
                connection.scalar(
                    sa.text("SELECT status FROM catalog_imports WHERE id=:id"),
                    {"id": raw_import_id},
                )
            )
            counts = (
                connection.execute(
                    sa.text(
                        "SELECT (SELECT COUNT(*) FROM product_styles WHERE site_id=1) "
                        "AS styles, (SELECT COUNT(*) FROM product_skus "
                        "WHERE site_id=1) "
                        "AS skus, (SELECT CAST(JSON_UNQUOTE(JSON_EXTRACT("
                        "effects.safe_result_json, '$.audit_only_row_count')) "
                        "AS UNSIGNED) FROM catalog_command_effects AS effects "
                        "INNER JOIN catalog_command_claims AS claims "
                        "ON claims.operation_token=effects.operation_token "
                        "WHERE claims.site_id=1 "
                        "AND claims.resource_kind='catalog_import' "
                        "AND claims.resource_id=:import_id "
                        "AND claims.operation_kind='approve' LIMIT 1) AS audit_only"
                    ),
                    {
                        "import_id": online_import_id,
                    },
                )
                .mappings()
                .one()
            )
    finally:
        engine.dispose()
    if revision != CURRENT_SCHEMA_REVISION:
        raise ValueError("live evidence migration head drifted")
    if int(row["current_id"]) != expected_version_id:
        raise ValueError("live result catalog version is not current")
    artifact_checksum = bytes(row["artifact_sha256"]).hex()
    if artifact_checksum != workbook_checksum:
        raise ValueError("downloaded workbook does not match the immutable export")
    if raw_status in {"uploaded", "parsing", "parse_failed"}:
        raise ValueError("raw workbook did not reach a reviewable state")
    if int(counts["styles"]) < 1 or int(counts["skus"]) < 1:
        raise ValueError("catalog approval did not create stable identities")
    if int(counts["audit_only"]) < 1:
        raise ValueError("live approval did not preserve an audit-only decision")
    evidence = {
        "run_id": run_id,
        "migration_head": revision,
        "catalog_version_id": expected_version_id,
        "catalog_version_number": int(row["version_number"]),
        "catalog_export_id": int(row["export_id"]),
        "workbook_sha256": workbook_checksum,
        "raw_import_status": raw_status,
        "stable_style_count": int(counts["styles"]),
        "stable_sku_count": int(counts["skus"]),
        "audit_only_row_count": int(counts["audit_only"]),
        "playwright_project": "catalog-milestone-a",
        "cleanup": "pending",
    }
    evidence_path.write_text(
        json.dumps(evidence, ensure_ascii=False, indent=2, sort_keys=True),
        encoding="utf-8",
    )
    manifest.state = "evidence_collected"
    manifest.details["catalog_version_id"] = expected_version_id
    manifest.details["workbook_sha256"] = workbook_checksum
    _save_manifest(manifest_path, manifest)
    return evidence


def _parse_path(value: str) -> Path:
    return Path(unquote(value)).resolve()


def main(arguments: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="product-catalog-live-harness")
    parser.add_argument("--run-id", required=True)
    parser.add_argument("--manifest", required=True)
    subcommands = parser.add_subparsers(dest="command", required=True)
    subcommands.add_parser("preflight")
    seed_parser = subcommands.add_parser("seed")
    seed_parser.add_argument("--include-catalog-experience", action="store_true")
    mutate = subcommands.add_parser("mutate-workbook")
    mutate.add_argument("--input", required=True)
    mutate.add_argument("--output", required=True)
    evidence = subcommands.add_parser("evidence")
    evidence.add_argument("--result", required=True)
    evidence.add_argument("--output", required=True)
    options = parser.parse_args(arguments)
    run_id = validate_run_id(options.run_id)
    manifest_path = _manifest_path(options.manifest, run_id=run_id)
    if options.command == "preflight":
        preflight(run_id=run_id, manifest_path=manifest_path)
    elif options.command == "seed":
        seed(
            run_id=run_id,
            manifest_path=manifest_path,
            include_catalog_experience=options.include_catalog_experience,
        )
    elif options.command == "mutate-workbook":
        mutate_workbook(
            run_id=run_id,
            manifest_path=manifest_path,
            input_path=_parse_path(options.input),
            output_path=_parse_path(options.output),
        )
    else:
        value = collect_evidence(
            run_id=run_id,
            manifest_path=manifest_path,
            result_path=_parse_path(options.result),
            evidence_path=_parse_path(options.output),
        )
        print(json.dumps(value, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())


__all__ = (
    "DatabaseIdentity",
    "ProcessIdentity",
    "RunManifest",
    "assert_database_identity",
    "assert_empty_seed_target",
    "assert_live_playwright_config",
    "assert_owned_cleanup_path",
    "assert_process_identity",
    "collect_evidence",
    "main",
    "mutate_workbook",
    "preflight",
    "seed",
    "validate_database_urls",
    "validate_run_id",
)
